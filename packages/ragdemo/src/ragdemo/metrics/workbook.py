"""口径 Excel 模板的导出与读取（总计划阶段 1 任务 1.3）。

三张可编辑的表 + 一张说明表。表头一行是机器列名（与数据库字段同名），中文含义、
是否必填与可选值写在「说明」表里。

读取是**宽松**的：单元格解析失败不立即抛出，而是记进 `WorkbookData.problems`，
由 `importer.plan_import` 与需要查库的校验（环节名、实体是否存在）合并后一次报全——
修一个错跑一次、再冒下一个，是录入者最讨厌的体验。
"""

from __future__ import annotations

import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Generic, TypeVar

import openpyxl
import psycopg
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ragdemo.metrics.definitions import (
    AiRevenueRuleSpec,
    MetricSpec,
    SourceMapSpec,
    current_ai_revenue_rules,
    current_metrics,
    current_source_maps,
)
from ragdemo.seed.taxonomy import DEFAULT_TAXONOMY_PATH, Taxonomy, load_taxonomy

METRIC_SHEET = "指标定义"
SOURCE_MAP_SHEET = "字段映射"
RULE_SHEET = "AI收入规则"
README_SHEET = "说明"

METRIC_ROLES = ("leading", "confirming")
CONFIDENCES = ("high", "medium", "low", "inferred")
METRIC_ID_PATTERN = re.compile(r"^[a-z][a-z0-9_]{2,63}$")

S = TypeVar("S")
T = TypeVar("T")


class WorkbookError(ValueError):
    """模板内容没通过校验。`problems` 是逐条的问题，每条都带表名与行号。"""

    def __init__(self, problems: Sequence[str]) -> None:
        self.problems = list(problems)
        super().__init__("口径模板没有通过校验：\n" + "\n".join(f"  - {p}" for p in problems))


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    required: bool
    note: str = ""


METRIC_COLUMNS = (
    Column(
        "metric_id",
        "指标 ID",
        True,
        "小写字母开头，只含小写字母/数字/下划线，3–64 位；是身份，不可改名",
    ),
    Column("metric_name", "指标名称", True),
    Column("metric_role", "指标角色", True, "leading（领先）/ confirming（确认）"),
    Column("frequency", "频率", True, "daily / monthly / quarterly / annual / irregular"),
    Column(
        "source_type", "来源类型", True, "filing / announcement / vendor_api / ir_page / derived"
    ),
    Column("definition", "口径定义", True, "这个数字具体是什么、怎么算，必填"),
    Column("unit", "单位", True, "CNY / USD / pcs / pct / ratio / x ……"),
    Column(
        "l3_node",
        "所属三级环节",
        False,
        "必须是 db/seed/taxonomy.csv 在册的环节名；跨环节指标留空",
    ),
    Column("source_entity", "来源实体", False, "多个用 | 分隔"),
    Column("extraction_hint", "抽取提示", False, "给抽取 Agent 的定位提示"),
    Column("direction", "方向", False, "1（越大越好）或 -1（越小越好），留空按 1"),
)
SOURCE_MAP_COLUMNS = (
    Column("metric_id", "指标 ID", True, "必须已存在，或在本文件「指标定义」表里新增"),
    Column("provider", "供应商", True, "如 tushare"),
    Column("provider_field", "供应商字段", True, "供应商返回里的字段名"),
    Column("priority", "优先级", False, "整数，越小越优先，留空按 100"),
    Column("scale_factor", "换算系数", False, "供应商单位 × 系数 = 本系统单位，留空按 1"),
    Column("notes", "备注", False),
)
RULE_COLUMNS = (
    Column("entity_id", "实体 ID", True, "必须是 core.entity 里已有的实体，如 CN.688041"),
    Column("rule_text", "拆分规则", True, "AI 相关收入怎么从总收入里拆出来"),
    Column(
        "source_type",
        "依据类型",
        True,
        "segment_report / management_guidance / analyst / estimate",
    ),
    Column("confidence", "置信度", True, "high / medium / low / inferred"),
    Column("effective_from", "生效日期", True, "YYYY-MM-DD；规则描述的是从哪天起的口径"),
)


@dataclass(frozen=True)
class Row(Generic[S]):
    sheet: str
    row_number: int
    spec: S


@dataclass
class WorkbookData:
    metrics: list[Row[MetricSpec]] = field(default_factory=list)
    source_maps: list[Row[SourceMapSpec]] = field(default_factory=list)
    rules: list[Row[AiRevenueRuleSpec]] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)


# --- 导出 ------------------------------------------------------------------------


def _cell(value: object) -> object:
    if isinstance(value, tuple):
        return "|".join(str(v) for v in value)
    if isinstance(value, Decimal):
        return str(value)
    return value


def _write_sheet(ws: Worksheet, columns: Sequence[Column], rows: Sequence[object]) -> None:
    ws.append([c.key for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for spec in rows:
        ws.append([_cell(getattr(spec, c.key)) for c in columns])
    ws.freeze_panes = "A2"
    for i, column in enumerate(columns, start=1):
        width = max(
            len(column.key), *(len(str(_cell(getattr(r, column.key)) or "")) for r in rows), 8
        )
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)


def export_template(conn: psycopg.Connection, path: Path) -> None:
    """导出当前生效的口径，作为编辑的起点。"""
    wb = openpyxl.Workbook()
    readme = wb.active
    assert readme is not None
    readme.title = README_SHEET
    readme.append(["表", "列名", "含义", "必填", "说明"])
    for cell in readme[1]:
        cell.font = Font(bold=True)
    for sheet, columns in (
        (METRIC_SHEET, METRIC_COLUMNS),
        (SOURCE_MAP_SHEET, SOURCE_MAP_COLUMNS),
        (RULE_SHEET, RULE_COLUMNS),
    ):
        for c in columns:
            readme.append([sheet, c.key, c.label, "是" if c.required else "", c.note])
    readme.append([])
    readme.append(["规则", "", "改一行 = 追加一个新版本，旧版本永久保留；删掉一行不会删除口径"])
    readme.append(
        ["规则", "", "导入默认只预览差异，确认无误后加 --yes 才写库；整批要么全成、要么全不写"]
    )
    readme.append(["规则", "", "第一行列名不要改，列的顺序可以调整"])

    metrics = sorted(current_metrics(conn).values(), key=lambda m: m.metric_id)
    maps = sorted(current_source_maps(conn).values(), key=lambda m: (m.metric_id, m.provider))
    rules = sorted(current_ai_revenue_rules(conn).values(), key=lambda r: r.entity_id)
    _write_sheet(wb.create_sheet(METRIC_SHEET), METRIC_COLUMNS, metrics)
    _write_sheet(wb.create_sheet(SOURCE_MAP_SHEET), SOURCE_MAP_COLUMNS, maps)
    _write_sheet(wb.create_sheet(RULE_SHEET), RULE_COLUMNS, rules)
    wb.save(path)


# --- 读取 ------------------------------------------------------------------------


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _int(value: object, *, default: int) -> int:
    text = _text(value)
    if text is None:
        return default
    number = float(text)
    if number != int(number):
        raise ValueError(f"应为整数，收到 {text!r}")
    return int(number)


def _decimal(value: object, *, default: Decimal) -> Decimal:
    text = _text(value)
    if text is None:
        return default
    try:
        return Decimal(text)
    except InvalidOperation:
        raise ValueError(f"应为数字，收到 {text!r}") from None


def _date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if text is None:
        raise ValueError("为空")
    return date.fromisoformat(text)


def _choice(value: object, choices: Sequence[str]) -> str:
    text = _text(value)
    if text is None or text not in choices:
        raise ValueError(f"应为 {' / '.join(choices)} 之一，收到 {text!r}")
    return text


class _Fields:
    """逐字段取值并收集错误：一行里有几处不对就报几处，而不是停在第一处。"""

    def __init__(self, get: Callable[[str], object]) -> None:
        self._get = get
        self.errors: list[str] = []

    def take(self, key: str, convert: Callable[[object], T]) -> T | None:
        try:
            return convert(self._get(key))
        except ValueError as exc:
            self.errors.append(f"{key}: {exc}")
            return None

    def done(self) -> None:
        if self.errors:
            raise ValueError("；".join(self.errors))


def _required_text(value: object) -> str:
    text = _text(value)
    if text is None:
        raise ValueError("必填，当前为空")
    return text


def _metric_id(value: object) -> str:
    text = _text(value) or ""
    if not METRIC_ID_PATTERN.match(text):
        raise ValueError(f"{text!r} 不合法（小写字母开头，只含 a-z0-9_，3–64 位）")
    return text


def _direction(value: object) -> int:
    direction = _int(value, default=1)
    if direction not in (-1, 1):
        raise ValueError(f"应为 1 或 -1，收到 {direction}")
    return direction


def _entities(value: object) -> tuple[str, ...]:
    text = _text(value)
    return tuple(s.strip() for s in text.split("|") if s.strip()) if text else ()


def _metric_parser(taxonomy: Taxonomy) -> Callable[[Callable[[str], object]], MetricSpec]:
    def l3_node(value: object) -> str | None:
        text = _text(value)
        if text and taxonomy.unknown([text]):
            raise ValueError(f"「{text}」不在 taxonomy.csv 环节表里")
        return text

    def parse(get: Callable[[str], object]) -> MetricSpec:
        f = _Fields(get)
        values = {
            "metric_id": f.take("metric_id", _metric_id),
            "metric_name": f.take("metric_name", _required_text),
            "metric_role": f.take("metric_role", lambda v: _choice(v, METRIC_ROLES)),
            "frequency": f.take("frequency", _required_text),
            "source_type": f.take("source_type", _required_text),
            "definition": f.take("definition", _required_text),
            "unit": f.take("unit", _required_text),
            "l3_node": f.take("l3_node", l3_node),
            "source_entity": f.take("source_entity", _entities),
            "extraction_hint": f.take("extraction_hint", _text),
            "direction": f.take("direction", _direction),
        }
        f.done()
        return MetricSpec(**values)  # type: ignore[arg-type]

    return parse


def _source_map(get: Callable[[str], object]) -> SourceMapSpec:
    f = _Fields(get)
    values = {
        "metric_id": f.take("metric_id", _required_text),
        "provider": f.take("provider", _required_text),
        "provider_field": f.take("provider_field", _required_text),
        "priority": f.take("priority", lambda v: _int(v, default=100)),
        "scale_factor": f.take("scale_factor", lambda v: _decimal(v, default=Decimal(1))),
        "notes": f.take("notes", _text),
    }
    f.done()
    return SourceMapSpec(**values)  # type: ignore[arg-type]


def _rule(get: Callable[[str], object]) -> AiRevenueRuleSpec:
    f = _Fields(get)
    values = {
        "entity_id": f.take("entity_id", _required_text),
        "rule_text": f.take("rule_text", _required_text),
        "source_type": f.take("source_type", _required_text),
        "confidence": f.take("confidence", lambda v: _choice(v, CONFIDENCES)),
        "effective_from": f.take("effective_from", _date),
    }
    f.done()
    return AiRevenueRuleSpec(**values)  # type: ignore[arg-type]


def _read_sheet(
    wb: openpyxl.Workbook,
    sheet: str,
    columns: Sequence[Column],
    parse: Callable[[Callable[[str], object]], S],
    problems: list[str],
) -> list[Row[S]]:
    if sheet not in wb.sheetnames:
        problems.append(f"缺少工作表「{sheet}」")
        return []
    ws = wb[sheet]
    header = [_text(c.value) for c in ws[1]]
    missing = [c.key for c in columns if c.key not in header]
    if missing:
        problems.append(f"{sheet}: 表头缺少列 {missing}")
        return []
    rows: list[Row[S]] = []
    for number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(_text(v) is None for v in values):
            continue
        cells: dict[str, Any] = {k: v for k, v in zip(header, values, strict=False) if k}

        def get(key: str, cells: dict[str, Any] = cells) -> object:
            return cells.get(key)

        try:
            rows.append(Row(sheet, number, parse(get)))
        except ValueError as exc:
            key = _text(cells.get(columns[0].key)) or "（空）"
            problems.append(f"{sheet} 第 {number} 行（{columns[0].key}={key}）: {exc}")
    return rows


def read_workbook(path: Path, *, taxonomy: Taxonomy | None = None) -> WorkbookData:
    tax = taxonomy or load_taxonomy(DEFAULT_TAXONOMY_PATH)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    data = WorkbookData()
    data.metrics = _read_sheet(wb, METRIC_SHEET, METRIC_COLUMNS, _metric_parser(tax), data.problems)
    data.source_maps = _read_sheet(
        wb, SOURCE_MAP_SHEET, SOURCE_MAP_COLUMNS, _source_map, data.problems
    )
    data.rules = _read_sheet(wb, RULE_SHEET, RULE_COLUMNS, _rule, data.problems)
    wb.close()
    return data
