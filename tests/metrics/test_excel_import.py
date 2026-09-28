"""口径 Excel 导入（总计划阶段 1 任务 1.3）。

流程：导出模板（预填当前生效口径）→ 在 Excel 里改 → 导入：先校验、再展示修改前后
差异，确认（--yes）后按版本追加写入，一整批放在一个事务里。
"""

from __future__ import annotations

from collections.abc import Iterator
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import psycopg
import pytest
from click.testing import CliRunner

from ragdemo.cli import main
from ragdemo.metrics.definitions import MetricSpec, PublishOutcome, metric_as_of, publish_metric
from ragdemo.metrics.importer import ImportRejected, apply_import, plan_import
from ragdemo.metrics.workbook import (
    METRIC_SHEET,
    RULE_SHEET,
    SOURCE_MAP_SHEET,
    WorkbookError,
    export_template,
    read_workbook,
)
from ragdemo.seed.loader import load_all
from ragdemo_core.db.migrate import migrate

MIGRATIONS = Path("db/migrations")
SEED = Path("db/seed")


@pytest.fixture
def conn(temp_db: str) -> Iterator[psycopg.Connection]:
    c = psycopg.connect(temp_db)
    migrate(c, MIGRATIONS)
    load_all(c, SEED)
    c.commit()
    yield c
    c.close()


def _set_cell(
    path: Path, sheet: str, key_col: str, key: str, col: str, value: str | int | None
) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[header.index(key_col)].value == key:
            row[header.index(col)].value = value
            break
    else:
        raise AssertionError(f"{sheet} 里没有 {key_col}={key}")
    wb.save(path)


def _append_row(path: Path, sheet: str, values: dict[str, str | int | None]) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    ws.append([values.get(str(h)) for h in header])
    wb.save(path)


def _versions(conn: psycopg.Connection, metric_id: str) -> int:
    row = conn.execute(
        "SELECT count(*) FROM core.node_metric_version WHERE metric_id = %s", (metric_id,)
    ).fetchone()
    assert row is not None
    return int(row[0])


@pytest.mark.db
def test_exported_template_round_trips_to_no_changes(
    conn: psycopg.Connection, tmp_path: Path
) -> None:
    path = tmp_path / "口径.xlsx"
    export_template(conn, path)
    plan = plan_import(conn, read_workbook(path))
    assert plan.changes and all(c.outcome is PublishOutcome.UNCHANGED for c in plan.changes)
    wb = openpyxl.load_workbook(path)
    assert {METRIC_SHEET, SOURCE_MAP_SHEET, RULE_SHEET, "说明"} <= set(wb.sheetnames)


@pytest.mark.db
def test_editing_one_definition_previews_one_diff_and_appends_a_version(
    conn: psycopg.Connection, tmp_path: Path
) -> None:
    path = tmp_path / "口径.xlsx"
    export_template(conn, path)
    old = metric_as_of(conn, "revenue_total", as_of=datetime.now().astimezone())
    assert old is not None
    new_definition = old.spec.definition + "（含其他业务收入）"
    _set_cell(path, METRIC_SHEET, "metric_id", "revenue_total", "definition", new_definition)

    plan = plan_import(conn, read_workbook(path))
    changed = [c for c in plan.changes if c.outcome is not PublishOutcome.UNCHANGED]
    assert len(changed) == 1
    (change,) = changed
    assert change.key == "revenue_total"
    assert change.outcome is PublishOutcome.NEW_VERSION
    assert change.diffs == [("definition", old.spec.definition, new_definition)]
    assert _versions(conn, "revenue_total") == 1, "预览不写库"

    imported_at = datetime.now().astimezone()
    counts = apply_import(conn, plan, author="founder", change_note="扩口径", known_at=imported_at)
    assert counts[PublishOutcome.NEW_VERSION] == 1
    assert _versions(conn, "revenue_total") == 2, "新增版本而不是覆盖"

    before = metric_as_of(conn, "revenue_total", as_of=imported_at - timedelta(seconds=1))
    after = metric_as_of(conn, "revenue_total", as_of=imported_at + timedelta(seconds=1))
    assert before is not None and after is not None
    assert before.spec.definition == old.spec.definition
    assert after.spec.definition == new_definition and after.author == "founder"


@pytest.mark.db
def test_all_validation_errors_are_reported_with_sheet_and_row(
    conn: psycopg.Connection, tmp_path: Path
) -> None:
    path = tmp_path / "口径.xlsx"
    export_template(conn, path)
    _set_cell(path, METRIC_SHEET, "metric_id", "revenue_total", "metric_role", "lagging")
    _set_cell(path, METRIC_SHEET, "metric_id", "rd_expense", "direction", 2)
    _set_cell(path, METRIC_SHEET, "metric_id", "rd_expense", "l3_node", "不存在的环节")
    _append_row(path, METRIC_SHEET, {"metric_id": "BadId", "metric_name": "x"})

    with pytest.raises(WorkbookError) as excinfo:
        plan_import(conn, read_workbook(path))
    message = str(excinfo.value)
    assert METRIC_SHEET in message
    for needle in ("lagging", "direction", "不存在的环节", "BadId"):
        assert needle in message, needle
    assert _versions(conn, "revenue_total") == 1


@pytest.mark.db
def test_source_map_may_reference_a_metric_added_in_the_same_workbook(
    conn: psycopg.Connection, tmp_path: Path
) -> None:
    path = tmp_path / "口径.xlsx"
    export_template(conn, path)
    _append_row(
        path,
        METRIC_SHEET,
        {
            "metric_id": "gross_margin",
            "metric_name": "毛利率",
            "metric_role": "confirming",
            "frequency": "quarterly",
            "source_type": "derived",
            "definition": "（营业收入 - 营业成本）/ 营业收入",
            "unit": "ratio",
        },
    )
    _append_row(
        path,
        SOURCE_MAP_SHEET,
        {
            "metric_id": "gross_margin",
            "provider": "tushare",
            "provider_field": "grossprofit_margin",
        },
    )
    _append_row(
        path,
        SOURCE_MAP_SHEET,
        {"metric_id": "no_such_metric", "provider": "tushare", "provider_field": "x"},
    )
    with pytest.raises(WorkbookError, match="no_such_metric"):
        plan_import(conn, read_workbook(path))


@pytest.mark.db
def test_rows_removed_from_the_sheet_are_left_alone(
    conn: psycopg.Connection, tmp_path: Path
) -> None:
    """表里删掉一行不等于删除口径：导入只追加/修改表里出现的行。"""
    path = tmp_path / "口径.xlsx"
    export_template(conn, path)
    before = conn.execute("SELECT count(*) FROM core.node_metric").fetchone()
    wb = openpyxl.load_workbook(path)
    wb[METRIC_SHEET].delete_rows(2)
    wb.save(path)
    plan = plan_import(conn, read_workbook(path))
    apply_import(
        conn, plan, author="founder", change_note=None, known_at=datetime.now().astimezone()
    )
    row = conn.execute("SELECT count(*) FROM core.node_metric").fetchone()
    assert row is not None and before is not None and row[0] == before[0]


@pytest.mark.db
def test_failure_mid_import_rolls_back_the_whole_batch(
    conn: psycopg.Connection, tmp_path: Path
) -> None:
    """rd_expense 已经有一个更晚的版本：批次里 revenue_total 先写成功、rd_expense
    随后被拒（乱序），整批必须回滚——revenue_total 不能留下半截修改。"""
    path = tmp_path / "口径.xlsx"
    export_template(conn, path)
    _set_cell(path, METRIC_SHEET, "metric_id", "revenue_total", "definition", "改一")
    _set_cell(path, METRIC_SHEET, "metric_id", "rd_expense", "definition", "改二")
    plan = plan_import(conn, read_workbook(path))

    now = datetime.now().astimezone()
    newer = metric_as_of(conn, "rd_expense", as_of=now)
    assert newer is not None
    publish_metric(
        conn,
        MetricSpec(**{**newer.spec.__dict__, "unit": "CNY_100M"}),
        author="other",
        known_at=now + timedelta(hours=1),
        ingest_run_id="t",
    )
    conn.commit()

    with pytest.raises(ImportRejected, match="rd_expense"):
        apply_import(conn, plan, author="founder", change_note=None, known_at=now)
    assert _versions(conn, "revenue_total") == 1
    assert _versions(conn, "rd_expense") == 2  # 只有那个更晚的版本，批次没留下任何东西


# --- CLI 端到端 ------------------------------------------------------------------


@pytest.mark.db
def test_cli_export_edit_dry_run_then_apply(
    conn: psycopg.Connection, temp_db: str, tmp_path: Path
) -> None:
    path = tmp_path / "口径.xlsx"
    env = {"RAGDEMO_DSN": temp_db}
    exported = CliRunner().invoke(main, ["metrics", "export-template", "--out", str(path)], env=env)
    assert exported.exit_code == 0, exported.output
    _set_cell(path, METRIC_SHEET, "metric_id", "revenue_total", "unit", "CNY_10K")

    dry = CliRunner().invoke(
        main, ["metrics", "import", "--file", str(path), "--author", "founder"], env=env
    )
    assert dry.exit_code == 0, dry.output
    assert "revenue_total" in dry.output and "CNY_10K" in dry.output
    assert "dry run" in dry.output
    assert _versions(conn, "revenue_total") == 1

    applied = CliRunner().invoke(
        main,
        [
            "metrics",
            "import",
            "--file",
            str(path),
            "--author",
            "founder",
            "--note",
            "改单位",
            "--yes",
        ],
        env=env,
    )
    assert applied.exit_code == 0, applied.output
    assert _versions(conn, "revenue_total") == 2


@pytest.mark.db
def test_cli_import_requires_an_author(temp_db: str, tmp_path: Path) -> None:
    result = CliRunner().invoke(
        main,
        ["metrics", "import", "--file", str(tmp_path / "x.xlsx")],
        env={"RAGDEMO_DSN": temp_db},
    )
    assert result.exit_code != 0
